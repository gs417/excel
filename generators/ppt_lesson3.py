"""PPT Lesson 3: 提案書作成（総合）
課題: 空テンプレ + シナリオ → 「提案書を作れ」
模範: 課題/提案/効果の3枚完成例
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

FONT_NAME = "游ゴシック"
BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x44, 0x72, 0xC4)
GREEN = RGBColor(0x70, 0xAD, 0x47)
ORANGE = RGBColor(0xED, 0x7D, 0x31)


def _add_textbox(slide, left, top, width, height, text, font_size=Pt(12),
                 bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.name = FONT_NAME
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return txBox


def _add_header_bar(slide, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()
    _add_textbox(slide, Inches(0.5), Inches(0.15), Inches(9), Inches(0.7),
                 text, Pt(24), True, WHITE)


def generate(output_dir: str):
    # ========== 課題ファイル ==========
    prs_q = Presentation()
    prs_q.slide_width = Inches(10)
    prs_q.slide_height = Inches(7.5)

    # 指示スライド
    slide0 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_textbox(slide0, Inches(0.5), Inches(1.5), Inches(9), Inches(1),
                 "課題: 以下のシナリオを基に提案書を完成させてください",
                 Pt(26), True, BLUE, PP_ALIGN.CENTER)

    scenario = (
        "【シナリオ】\n"
        "あなたはミルクネクスト株式会社の法人営業担当です。\n"
        "取引先: グリーンバスケット（スーパー、関東50店舗展開）\n\n"
        "【取引先の状況】\n"
        "• 乳製品の売上が前年比95%で伸び悩んでいる\n"
        "• 現在の乳製品コーナーは棚2本分で、競合品中心の品揃え\n"
        "• 健康志向商品の顧客ニーズが高まっているが対応できていない\n"
        "• 週末の集客施策を探しており、試食イベントに関心あり\n\n"
        "【提案内容】\n"
        "ギリシャ濃密ヨーグルト + 北海道フレッシュ牛乳 の重点導入提案\n"
        "• 棚割り: 乳製品棚3本に拡大（当社商品で1本分確保）\n"
        "• 販促: 毎月第1土曜に試食販売（当社販売員を派遣）\n"
        "• 期待効果: 乳製品カテゴリ売上15%アップ（月額約120万円の売上増）\n"
        "• リベート: 導入月に仕入額の5%（約50万円）を販促協力金として提供\n\n"
        "【指示】\n"
        "次の3枚のスライドテンプレートに内容を記入してください:\n"
        "1. 現状の課題（3つの課題を整理）\n"
        "2. ソリューション提案（何をどう解決するか）\n"
        "3. 導入効果とスケジュール（数値根拠と導入計画）"
    )
    _add_textbox(slide0, Inches(0.5), Inches(2.5), Inches(9), Inches(4.5),
                 scenario, Pt(11), False, DARK_GRAY)

    # テンプレートスライド1: 課題
    slide_t1 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_header_bar(slide_t1, "1. 現状の課題")
    _add_textbox(slide_t1, Inches(0.5), Inches(1.5), Inches(9), Inches(0.5),
                 "グリーンバスケットが抱える3つの課題を整理してください", Pt(12), False, DARK_GRAY)

    for i in range(3):
        y = Inches(2.5) + Emu(int(Inches(1.5) * i))
        shape = slide_t1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                          Inches(0.8), y, Inches(8.4), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"課題 {i + 1}: ここに記入"
        p.font.size = Pt(14)
        p.font.name = FONT_NAME
        p.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    # テンプレートスライド2: 提案
    slide_t2 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_header_bar(slide_t2, "2. ソリューション提案")
    _add_textbox(slide_t2, Inches(0.5), Inches(1.5), Inches(9), Inches(0.5),
                 "棚割り提案と販促施策でどう解決するか記入してください",
                 Pt(12), False, DARK_GRAY)

    for i, label in enumerate(["棚割り提案の内容", "販促施策の内容"]):
        x = Inches(0.5) + Emu(int(Inches(4.7) * i))
        shape = slide_t2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                          x, Inches(2.5), Inches(4.3), Inches(4))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{label}\n\nここに記入"
        p.font.size = Pt(12)
        p.font.name = FONT_NAME
        p.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    # テンプレートスライド3: 効果
    slide_t3 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_header_bar(slide_t3, "3. 導入効果とスケジュール")
    _add_textbox(slide_t3, Inches(0.5), Inches(1.5), Inches(9), Inches(0.5),
                 "数値根拠と導入スケジュールを記入してください",
                 Pt(12), False, DARK_GRAY)

    for i, label in enumerate(["期待効果（数値根拠）", "導入スケジュール"]):
        y = Inches(2.5) + Emu(int(Inches(2.2) * i))
        shape = slide_t3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                          Inches(0.8), y, Inches(8.4), Inches(1.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{label}\n\nここに記入"
        p.font.size = Pt(12)
        p.font.name = FONT_NAME
        p.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    path_q = f"{output_dir}/ppt/Lesson3_課題_提案書作成.pptx"
    prs_q.save(path_q)
    print(f"  ✓ {path_q}")

    # ========== 模範解答ファイル ==========
    prs_a = Presentation()
    prs_a.slide_width = Inches(10)
    prs_a.slide_height = Inches(7.5)

    # 表紙
    slide_cover = prs_a.slides.add_slide(prs_a.slide_layouts[6])
    bg = slide_cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE
    bg.line.fill.background()
    _add_textbox(slide_cover, Inches(1), Inches(2), Inches(8), Inches(1.5),
                 "乳製品売場活性化のご提案", Pt(32), True, WHITE, PP_ALIGN.CENTER)
    _add_textbox(slide_cover, Inches(1), Inches(4), Inches(8), Inches(1),
                 "グリーンバスケット 御中", Pt(18), False, WHITE, PP_ALIGN.CENTER)
    _add_textbox(slide_cover, Inches(1), Inches(5.5), Inches(8), Inches(1),
                 "ミルクネクスト株式会社 法人営業部", Pt(14), False, WHITE, PP_ALIGN.CENTER)

    # スライド1: 課題
    slide1 = prs_a.slides.add_slide(prs_a.slide_layouts[6])
    _add_header_bar(slide1, "1. 現状の課題")

    issues = [
        ("乳製品売上の伸び悩み", "前年比95%で低下傾向\n→ カテゴリ全体の活性化が必要", ORANGE),
        ("棚スペースの不足", "棚2本分と限定的、競合品中心の品揃え\n→ 健康志向商品が不足", ORANGE),
        ("集客施策の欠如", "週末の来店動機が弱い\n→ 体験型施策（試食等）で来店頻度向上の余地", RGBColor(0xC0, 0x00, 0x00)),
    ]

    for i, (title, detail, color) in enumerate(issues):
        y = Inches(1.5) + Emu(int(Inches(1.8) * i))
        circle = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), y, Inches(0.6), Inches(0.6))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        tf = circle.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].text = str(i + 1)
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].font.name = FONT_NAME

        _add_textbox(slide1, Inches(1.4), y, Inches(3), Inches(0.5),
                     title, Pt(16), True, DARK_GRAY)
        _add_textbox(slide1, Inches(1.4), Emu(int(y + Inches(0.5))), Inches(7.5), Inches(1),
                     detail, Pt(11), False, DARK_GRAY)

    # スライド2: 提案
    slide2 = prs_a.slides.add_slide(prs_a.slide_layouts[6])
    _add_header_bar(slide2, "2. ソリューション提案")

    box1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(0.3), Inches(1.5), Inches(4.5), Inches(5))
    box1.fill.solid()
    box1.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFE)
    box1.line.color.rgb = ACCENT
    box1.line.width = Pt(2)

    _add_textbox(slide2, Inches(0.5), Inches(1.7), Inches(4), Inches(0.5),
                 "棚割りリニューアル", Pt(18), True, ACCENT, PP_ALIGN.CENTER)
    cb_details = (
        "• 乳製品棚を2本→3本に拡大\n"
        "• 当社商品で1本分を確保\n"
        "• 北海道フレッシュ牛乳をゴールデンゾーンに配置\n"
        "• ギリシャ濃密ヨーグルトを健康訴求POPと共に展開"
    )
    _add_textbox(slide2, Inches(0.6), Inches(2.5), Inches(4), Inches(3),
                 cb_details, Pt(12), False, DARK_GRAY)

    box2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(5.2), Inches(1.5), Inches(4.5), Inches(5))
    box2.fill.solid()
    box2.fill.fore_color.rgb = RGBColor(0xE8, 0xF8, 0xE8)
    box2.line.color.rgb = GREEN
    box2.line.width = Pt(2)

    _add_textbox(slide2, Inches(5.4), Inches(1.7), Inches(4), Inches(0.5),
                 "販促施策", Pt(18), True, GREEN, PP_ALIGN.CENTER)
    di_details = (
        "• 毎月第1土曜に試食販売イベント実施\n"
        "• 当社から販売スタッフを2名派遣（費用当社負担）\n"
        "• 季節限定フレーバーの先行販売権を提供\n"
        "• SNS連動キャンペーン（#朝ミルク習慣）"
    )
    _add_textbox(slide2, Inches(5.5), Inches(2.5), Inches(4), Inches(3),
                 di_details, Pt(12), False, DARK_GRAY)

    _add_textbox(slide2, Inches(4.5), Inches(3.5), Inches(1), Inches(1),
                 "+", Pt(36), True, DARK_GRAY, PP_ALIGN.CENTER)

    # スライド3: 効果
    slide3 = prs_a.slides.add_slide(prs_a.slide_layouts[6])
    _add_header_bar(slide3, "3. 導入効果とスケジュール")

    _add_textbox(slide3, Inches(0.5), Inches(1.3), Inches(4), Inches(0.5),
                 "期待効果", Pt(18), True, BLUE)

    roi_items = [
        ("乳製品カテゴリ売上", "月額800万→920万（15%アップ）"),
        ("年間売上増加額", "約1,440万円"),
        ("導入月リベート", "仕入額の5%（約50万円）"),
        ("投資回収", "棚変更費用は販促協力金で初月回収"),
    ]
    for i, (label, value) in enumerate(roi_items):
        y = Inches(2) + Emu(int(Inches(0.5) * i))
        _add_textbox(slide3, Inches(0.8), y, Inches(2), Inches(0.4),
                     label, Pt(11), True, DARK_GRAY)
        _add_textbox(slide3, Inches(3), y, Inches(4), Inches(0.4),
                     value, Pt(11), False, DARK_GRAY)

    _add_textbox(slide3, Inches(0.5), Inches(4.3), Inches(4), Inches(0.5),
                 "導入スケジュール", Pt(18), True, BLUE)

    phases = [
        ("Month 1", "棚割り変更", "新レイアウト設計\n什器手配・POP制作", ACCENT),
        ("Month 2", "導入・試食", "商品入替え・棚設置\n初回試食イベント開催", GREEN),
        ("Month 3", "効果検証", "売上データ分析\n次月の改善提案", ORANGE),
    ]

    for i, (month, phase, detail, color) in enumerate(phases):
        x = Inches(0.5) + Emu(int(Inches(3) * i))
        shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                        x, Inches(5), Inches(2.8), Inches(1.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = f"{month}\n{phase}"
        p.font.size = Pt(14)
        p.font.name = FONT_NAME
        p.font.bold = True
        p.font.color.rgb = WHITE
        p2 = tf.add_paragraph()
        p2.text = detail
        p2.font.size = Pt(9)
        p2.font.name = FONT_NAME
        p2.font.color.rgb = WHITE
        p2.alignment = PP_ALIGN.CENTER

        if i < 2:
            arrow = slide3.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Emu(int(x + Inches(2.8))), Inches(5.5),
                Inches(0.2), Inches(0.5)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = DARK_GRAY
            arrow.line.fill.background()

    path_a = f"{output_dir}/ppt/Lesson3_模範_提案書作成.pptx"
    prs_a.save(path_a)
    print(f"  ✓ {path_a}")

    return [path_q, path_a]
