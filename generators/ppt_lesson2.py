"""PPT Lesson 2: 可視化（箇条書き → 図解）
課題: 箇条書きだけの2枚 → 「図解化せよ」
模範: ステップ矢印図 + プロポーショナルバー
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

FONT_NAME = "游ゴシック"
BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x44, 0x72, 0xC4)
COLORS = [
    RGBColor(0x44, 0x72, 0xC4),
    RGBColor(0xED, 0x7D, 0x31),
    RGBColor(0x70, 0xAD, 0x47),
    RGBColor(0xFF, 0xC0, 0x00),
    RGBColor(0xA5, 0xA5, 0xA5),
]


def _add_textbox(slide, left, top, width, height, text, font_size=Pt(12),
                 bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.name = FONT_NAME
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return txBox


def generate(output_dir: str):
    # ========== 課題ファイル ==========
    prs_q = Presentation()
    prs_q.slide_width = Inches(10)
    prs_q.slide_height = Inches(7.5)

    slide0 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_textbox(slide0, Inches(0.8), Inches(2.5), Inches(8.4), Inches(2),
                 "課題: 以下のスライドを図解に変換してください",
                 Pt(28), True, BLUE, PP_ALIGN.CENTER)
    _add_textbox(slide0, Inches(0.8), Inches(4.2), Inches(8.4), Inches(1),
                 "箇条書きを、矢印図・プロセス図・バーチャートなどのビジュアルに変換しましょう",
                 Pt(14), False, DARK_GRAY, PP_ALIGN.CENTER)

    # 課題スライド1: 営業プロセス
    slide1 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_textbox(slide1, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "新規取引先への営業プロセス", Pt(24), True, BLUE)

    process_text = (
        "• ステップ1: ターゲット選定 — 地域のスーパー・コンビニ・飲食チェーンの"
        "リストアップ。商圏分析で出店密度の高いエリアを優先。月間目標は30件リストアップ。\n\n"
        "• ステップ2: 初回訪問・ヒアリング — バイヤーとの面談で現在の仕入先・"
        "取扱カテゴリ・棚割り状況をヒアリング。競合品との価格差も確認。\n\n"
        "• ステップ3: 試食・サンプル提案 — 自社商品のサンプルを提供し、"
        "品質の優位性を体感してもらう。特にヨーグルト・チーズは試食が効果的。\n\n"
        "• ステップ4: 棚割り・条件交渉 — 導入時のリベート・棚位置・"
        "販促協力金を交渉。エンド陳列の獲得は売上に直結するため重点交渉。\n\n"
        "• ステップ5: 納品開始・フォロー — 初回納品後、2週間で売場チェック。"
        "売れ行きデータを基にリピート発注と棚の拡大を提案する。"
    )
    txBox = slide1.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = process_text
    p.font.size = Pt(11)
    p.font.name = FONT_NAME
    p.font.color.rgb = DARK_GRAY

    # 課題スライド2: 売上構成比
    slide2 = prs_q.slides.add_slide(prs_q.slide_layouts[6])
    _add_textbox(slide2, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "2025年度 カテゴリ別売上構成", Pt(24), True, BLUE)

    composition_text = (
        "• 牛乳: 売上全体の35%を占める主力。北海道フレッシュ牛乳1Lと濃厚4.2牛乳1Lが"
        "二本柱。前年比112%で安定成長。スーパーの定番棚を確保しており安定的。\n\n"
        "• ヨーグルト: 売上全体の28%。ギリシャ濃密ヨーグルトが健康志向層に好評で、"
        "特にコンビニチャネルでの伸びが顕著。PB受注の拡大余地あり。\n\n"
        "• チーズ: 売上全体の22%。熟成カマンベールが飲食チェーンで評価が高い。"
        "秋冬の鍋シーズンに向けてスライスチーズの需要増が見込まれる。\n\n"
        "• バター・生クリーム: 売上全体の15%。業務用の大容量規格が"
        "ホテル・洋菓子店に安定供給。北海道バターは原料高で利益率改善が課題。"
    )
    txBox2 = slide2.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = composition_text
    p2.font.size = Pt(11)
    p2.font.name = FONT_NAME
    p2.font.color.rgb = DARK_GRAY

    path_q = f"{output_dir}/ppt/Lesson2_課題_図解化.pptx"
    prs_q.save(path_q)
    print(f"  ✓ {path_q}")

    # ========== 模範解答ファイル ==========
    prs_a = Presentation()
    prs_a.slide_width = Inches(10)
    prs_a.slide_height = Inches(7.5)

    # --- 模範1: ステップ矢印図 ---
    slide_a1 = prs_a.slides.add_slide(prs_a.slide_layouts[6])

    header = slide_a1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1))
    header.fill.solid()
    header.fill.fore_color.rgb = BLUE
    header.line.fill.background()
    _add_textbox(slide_a1, Inches(0.5), Inches(0.15), Inches(9), Inches(0.7),
                 "新規取引先への営業プロセス", Pt(24), True, WHITE)

    steps = [
        ("1. ターゲット\n選定", "商圏分析\nリスト30件/月"),
        ("2. 初回訪問\nヒアリング", "バイヤー面談\n仕入先・棚状況"),
        ("3. 試食\nサンプル", "品質を体感\nヨーグルト・チーズ"),
        ("4. 棚割り\n条件交渉", "リベート・棚位置\nエンド陳列獲得"),
        ("5. 納品\nフォロー", "2週間後売場確認\nリピート提案"),
    ]

    box_w = Inches(1.5)
    box_h = Inches(1.2)
    arrow_w = Inches(0.3)
    start_x = Inches(0.5)
    y_top = Inches(1.8)
    y_detail = Inches(3.5)
    spacing = Inches(1.8)

    for i, (title, detail) in enumerate(steps):
        x = start_x + Emu(int(spacing * i))

        shape = slide_a1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_top, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS[i]
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.name = FONT_NAME
        p.font.bold = True
        p.font.color.rgb = WHITE

        _add_textbox(slide_a1, x, y_detail, box_w, Inches(2),
                     detail, Pt(9), False, DARK_GRAY, PP_ALIGN.CENTER)

        if i < 4:
            arrow_x = Emu(int(x + box_w))
            arrow = slide_a1.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, arrow_x, Emu(int(y_top + box_h / 3)),
                arrow_w, Inches(0.4)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = DARK_GRAY
            arrow.line.fill.background()

    _add_textbox(slide_a1, Inches(0.5), Inches(6), Inches(9), Inches(0.8),
                 "※ 各ステップのKPIを設定し、パイプライン管理を徹底する",
                 Pt(11), False, DARK_GRAY, PP_ALIGN.LEFT)

    # --- 模範2: プロポーショナルバー ---
    slide_a2 = prs_a.slides.add_slide(prs_a.slide_layouts[6])

    header2 = slide_a2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1))
    header2.fill.solid()
    header2.fill.fore_color.rgb = BLUE
    header2.line.fill.background()
    _add_textbox(slide_a2, Inches(0.5), Inches(0.15), Inches(9), Inches(0.7),
                 "2025年度 カテゴリ別売上構成", Pt(24), True, WHITE)

    categories = [
        ("牛乳", 35, "北海道フレッシュ/濃厚4.2\n前年比112%安定成長", COLORS[0]),
        ("ヨーグルト", 28, "ギリシャ濃密が健康志向に好評\nコンビニPB受注に余地", COLORS[1]),
        ("チーズ", 22, "熟成カマンベール飲食チェーンで好評\n秋冬鍋シーズンに需要増", COLORS[2]),
        ("バター・生クリーム", 15, "業務用ホテル・洋菓子店に安定\n原料高で利益率改善が課題", COLORS[3]),
    ]

    bar_y = Inches(1.5)
    bar_total_w = Inches(8.5)
    bar_h = Inches(1)
    bar_x_start = Inches(0.75)

    x_offset = bar_x_start
    for name, pct, detail, color in categories:
        w = Emu(int(bar_total_w * pct / 100))
        shape = slide_a2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_offset, bar_y, w, bar_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = f"{name}\n{pct}%"
        p.font.size = Pt(11)
        p.font.name = FONT_NAME
        p.font.bold = True
        p.font.color.rgb = WHITE
        x_offset = Emu(int(x_offset + w))

    card_y = Inches(3.2)
    card_w = Inches(4)
    card_h = Inches(1.5)
    positions = [
        (Inches(0.5), card_y),
        (Inches(5), card_y),
        (Inches(0.5), Inches(5)),
        (Inches(5), Inches(5)),
    ]

    for i, (name, pct, detail, color) in enumerate(categories):
        x, y = positions[i]
        indicator = slide_a2.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, y, Inches(0.15), card_h
        )
        indicator.fill.solid()
        indicator.fill.fore_color.rgb = color
        indicator.line.fill.background()

        _add_textbox(slide_a2, Emu(int(x + Inches(0.3))), y, Inches(3.5), Inches(0.5),
                     f"{name} ({pct}%)", Pt(14), True, DARK_GRAY)
        _add_textbox(slide_a2, Emu(int(x + Inches(0.3))), Emu(int(y + Inches(0.5))),
                     Inches(3.5), Inches(1),
                     detail, Pt(10), False, DARK_GRAY)

    path_a = f"{output_dir}/ppt/Lesson2_模範_図解化.pptx"
    prs_a.save(path_a)
    print(f"  ✓ {path_a}")

    return [path_q, path_a]
