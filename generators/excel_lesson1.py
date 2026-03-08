"""Excel Lesson 1: 基本関数 (SUM/AVERAGE/COUNT/COUNTIF/SUMIF) — 10問"""

import openpyxl
from openpyxl.utils import get_column_letter
from . import data, styles


def generate(output_dir: str):
    wb = openpyxl.Workbook()

    # ========== データシート ==========
    ws_data = wb.active
    ws_data.title = "売上データ"
    ws_data.sheet_properties.tabColor = "70AD47"

    headers = ["売上ID", "日付", "顧客ID", "商品名", "カテゴリ", "単価", "数量", "金額", "担当者", "地域"]
    for col, h in enumerate(headers, 1):
        ws_data.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_data, 1, 1, len(headers))

    for i, s in enumerate(data.SALES, 2):
        row_data = [
            s["id"], s["date"], s["customer_id"], s["product_name"],
            s["category"], s["unit_price"], s["quantity"], s["amount"],
            s["rep"], s["region"],
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws_data.cell(row=i, column=col, value=val)
            cell.font = styles.FONT_NORMAL
            cell.border = styles.THIN_BORDER
            if col == 2:
                cell.number_format = "YYYY/MM/DD"
            elif col in (6, 8):
                cell.number_format = styles.NUMBER_FORMAT_YEN

    # 列幅調整
    col_widths = [10, 12, 10, 24, 14, 14, 8, 14, 14, 10]
    for i, w in enumerate(col_widths, 1):
        ws_data.column_dimensions[get_column_letter(i)].width = w

    # ========== 問題シート ==========
    ws = wb.create_sheet("問題")
    styles.setup_problem_sheet(ws, "Lesson 1: 基本関数 (SUM / AVERAGE / COUNT / COUNTIF / SUMIF)")

    ws.cell(row=3, column=1, value="「売上データ」シートを参照して、以下の問題に回答してください。").font = styles.FONT_NORMAL

    # 期待値の計算
    amounts = [s["amount"] for s in data.SALES]
    quantities = [s["quantity"] for s in data.SALES]
    total_amount = sum(amounts)
    avg_amount = total_amount / len(amounts)
    total_count = len(data.SALES)
    tokyo_count = sum(1 for s in data.SALES if s["region"] == "東京")
    milk_amount = sum(s["amount"] for s in data.SALES if s["category"] == "牛乳")
    max_amount = max(amounts)
    min_amount = min(amounts)
    total_qty = sum(quantities)
    tanaka_count = sum(1 for s in data.SALES if s["rep"] == "田中 太郎")
    cheese_amount = sum(s["amount"] for s in data.SALES if s["category"] == "チーズ")

    questions = [
        ("Q1", "全売上の合計金額を求めてください（SUM関数）", "=SUM(売上データ!H2:H201)", total_amount),
        ("Q2", "売上金額の平均を求めてください（AVERAGE関数）", "=AVERAGE(売上データ!H2:H201)", avg_amount),
        ("Q3", "売上データの件数を求めてください（COUNT関数）", "=COUNT(売上データ!A2:A201)", total_count),
        ("Q4", "地域が「東京」の件数を求めてください（COUNTIF関数）", '=COUNTIF(売上データ!J2:J201,"東京")', tokyo_count),
        ("Q5", "カテゴリが「牛乳」の売上合計を求めてください（SUMIF関数）", '=SUMIF(売上データ!E2:E201,"牛乳",売上データ!H2:H201)', milk_amount),
        ("Q6", "売上金額の最大値を求めてください（MAX関数）", "=MAX(売上データ!H2:H201)", max_amount),
        ("Q7", "売上金額の最小値を求めてください（MIN関数）", "=MIN(売上データ!H2:H201)", min_amount),
        ("Q8", "全売上の数量（ケース数）合計を求めてください（SUM関数）", "=SUM(売上データ!G2:G201)", total_qty),
        ("Q9", "担当者「田中 太郎」の売上件数を求めてください（COUNTIF関数）", '=COUNTIF(売上データ!I2:I201,"田中 太郎")', tanaka_count),
        ("Q10", "カテゴリが「チーズ」の売上合計を求めてください（SUMIF関数）", '=SUMIF(売上データ!E2:E201,"チーズ",売上データ!H2:H201)', cheese_amount),
    ]

    # ヘッダー
    row = 5
    for col, h in enumerate(["No.", "問題", "ヒント", "あなたの回答", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 5)

    # 問題行
    answer_col = 4  # D列
    grading_col = 5  # E列
    start_row = row + 1

    hints = [
        "=SUM(範囲)", "=AVERAGE(範囲)", "=COUNT(範囲)",
        '=COUNTIF(範囲,"条件")', '=SUMIF(条件範囲,"条件",合計範囲)',
        "=MAX(範囲)", "=MIN(範囲)", "=SUM(範囲)",
        '=COUNTIF(範囲,"条件")', '=SUMIF(条件範囲,"条件",合計範囲)',
    ]

    for i, (qno, question, formula, expected) in enumerate(questions):
        r = start_row + i
        ws.cell(row=r, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=r, column=1).border = styles.THIN_BORDER
        ws.cell(row=r, column=2, value=question).font = styles.FONT_NORMAL
        ws.cell(row=r, column=2).border = styles.THIN_BORDER
        ws.cell(row=r, column=3, value=hints[i]).font = styles.FONT_NORMAL
        ws.cell(row=r, column=3).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, r, answer_col)
        answer_ref = f"D{r}"
        styles.add_grading_cell(ws, r, grading_col, answer_ref, expected)

    end_row = start_row + len(questions) - 1

    # 条件付き書式
    styles.apply_grading_format(ws, grading_col, start_row, end_row)

    # スコアサマリー
    styles.add_score_summary(ws, end_row + 2, 1, grading_col, start_row, end_row, len(questions))

    # 採点確認の説明
    ws.cell(row=end_row + 4, column=1,
            value="💡 全問回答後、E列のフォント色を黒に変更すると採点結果が見えます。").font = styles.FONT_NORMAL

    # 列幅
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 14

    # ========== 解答シート (非表示) ==========
    ws_ans = styles.create_answer_sheet(wb, ws)
    ws_ans.cell(row=1, column=1, value="No.").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=2, value="模範解答（数式）").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=3, value="期待値").font = styles.FONT_BOLD
    for i, (qno, _, formula, expected) in enumerate(questions):
        ws_ans.cell(row=i + 2, column=1, value=qno)
        ws_ans.cell(row=i + 2, column=2, value=formula)
        ws_ans.cell(row=i + 2, column=3, value=expected)

    # Move 問題 to first position
    wb.move_sheet("問題", offset=-1)

    path = f"{output_dir}/excel/Lesson1_基本関数.xlsx"
    wb.save(path)
    print(f"  ✓ {path}")
    return path
