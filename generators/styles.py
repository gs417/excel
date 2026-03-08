"""Excel共通スタイル・採点ヘルパー"""

from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle, numbers
)
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

# --- Colors ---
FILL_YELLOW = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")  # 回答セル
FILL_GREEN = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")   # 正解
FILL_RED = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")     # 不正解
FILL_HEADER = PatternFill(start_color="FF4472C4", end_color="FF4472C4", fill_type="solid")  # ヘッダー
FILL_LIGHT_BLUE = PatternFill(start_color="FFD6E4F0", end_color="FFD6E4F0", fill_type="solid")
FILL_LIGHT_GRAY = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")

FONT_HEADER = Font(name="游ゴシック", size=11, bold=True, color="FFFFFF")
FONT_NORMAL = Font(name="游ゴシック", size=11)
FONT_BOLD = Font(name="游ゴシック", size=11, bold=True)
FONT_TITLE = Font(name="游ゴシック", size=14, bold=True)
FONT_HIDDEN = Font(name="游ゴシック", size=11, color="FFFFFF")  # 白字（隠蔽用）
FONT_CORRECT = Font(name="游ゴシック", size=11, color="006100")
FONT_WRONG = Font(name="游ゴシック", size=11, color="9C0006")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

NUMBER_FORMAT_YEN = '#,##0"円"'
NUMBER_FORMAT_INT = '#,##0'
NUMBER_FORMAT_PCT = '0.0%'


def style_header_row(ws, row, col_start, col_end):
    """ヘッダー行にスタイル適用"""
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER


def style_data_cell(ws, row, col, font=None, fill=None, alignment=None, number_format=None):
    """データセルにスタイル適用"""
    cell = ws.cell(row=row, column=col)
    cell.font = font or FONT_NORMAL
    if fill:
        cell.fill = fill
    cell.alignment = alignment or ALIGN_CENTER
    cell.border = THIN_BORDER
    if number_format:
        cell.number_format = number_format
    return cell


def mark_answer_cell(ws, row, col):
    """回答セルに黄色背景を適用"""
    cell = ws.cell(row=row, column=col)
    cell.fill = FILL_YELLOW
    cell.font = FONT_BOLD
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER
    return cell


def add_grading_cell(ws, row, grading_col, answer_cell_ref, expected_value,
                     is_numeric=True, tolerance=0.01):
    """採点セル追加（初期状態は白字で隠蔽）

    Args:
        ws: worksheet
        row: 行番号
        grading_col: 採点結果を書く列番号
        answer_cell_ref: 回答セルの参照（例: "E5"）
        expected_value: 期待値（数値 or 文字列）
        is_numeric: 数値比較かどうか
        tolerance: 数値比較の許容誤差（割合）
    """
    cell = ws.cell(row=row, column=grading_col)
    if is_numeric:
        # 数値の場合: 誤差許容
        cell.value = (
            f'=IF({answer_cell_ref}="","",IF(ABS({answer_cell_ref}-{expected_value})'
            f'<ABS({expected_value}*{tolerance}),"✓ 正解","✗ 不正解"))'
        )
    else:
        # 文字列の場合: 完全一致
        cell.value = (
            f'=IF({answer_cell_ref}="","",IF({answer_cell_ref}={expected_value},"✓ 正解","✗ 不正解"))'
        )
    cell.font = FONT_HIDDEN  # 白字で隠蔽
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER
    return cell


def add_grading_formula_cell(ws, row, grading_col, answer_cell_ref, expected_cell_ref):
    """セル参照同士の比較による採点セル"""
    cell = ws.cell(row=row, column=grading_col)
    cell.value = (
        f'=IF({answer_cell_ref}="","",IF(ABS({answer_cell_ref}-{expected_cell_ref})'
        f'<ABS({expected_cell_ref}*0.01),"✓ 正解","✗ 不正解"))'
    )
    cell.font = FONT_HIDDEN
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER
    return cell


def apply_grading_format(ws, grading_col, start_row, end_row):
    """採点列に条件付き書式（正解=緑、不正解=赤）を適用"""
    col_letter = get_column_letter(grading_col)
    range_str = f"{col_letter}{start_row}:{col_letter}{end_row}"

    ws.conditional_formatting.add(
        range_str,
        CellIsRule(
            operator="containsText",
            formula=['"✓"'],
            font=FONT_CORRECT,
            fill=FILL_GREEN,
        ),
    )
    ws.conditional_formatting.add(
        range_str,
        CellIsRule(
            operator="containsText",
            formula=['"✗"'],
            font=FONT_WRONG,
            fill=FILL_RED,
        ),
    )


def add_score_summary(ws, row, col, grading_col, start_row, end_row, total_questions):
    """スコアサマリーを追加"""
    col_letter = get_column_letter(grading_col)
    range_str = f"{col_letter}{start_row}:{col_letter}{end_row}"

    ws.cell(row=row, column=col, value="正答数:").font = FONT_BOLD
    score_cell = ws.cell(row=row, column=col + 1)
    score_cell.value = f'=COUNTIF({range_str},"✓ 正解")'
    score_cell.font = FONT_BOLD

    ws.cell(row=row, column=col + 2, value=f"/ {total_questions} 問").font = FONT_BOLD

    ws.cell(row=row + 1, column=col, value="正答率:").font = FONT_BOLD
    rate_cell = ws.cell(row=row + 1, column=col + 1)
    rate_cell.value = f'=IF({total_questions}=0,0,COUNTIF({range_str},"✓ 正解")/{total_questions})'
    rate_cell.number_format = NUMBER_FORMAT_PCT
    rate_cell.font = FONT_BOLD


def setup_problem_sheet(ws, title):
    """問題シートの共通セットアップ"""
    ws.sheet_properties.tabColor = "4472C4"
    ws.cell(row=1, column=1, value=title).font = FONT_TITLE
    ws.cell(row=2, column=1, value="黄色のセルに回答を入力してください。").font = FONT_NORMAL
    ws.cell(row=2, column=1).fill = FILL_YELLOW


def create_answer_sheet(wb, ws_problem, name="解答"):
    """非表示の解答シートを作成"""
    ws = wb.create_sheet(name)
    ws.sheet_state = "hidden"
    return ws
