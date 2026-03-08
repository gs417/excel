"""Excel Lesson 2: VLOOKUP / INDEX-MATCH — 8問"""

import openpyxl
from openpyxl.utils import get_column_letter
from . import data, styles


def generate(output_dir: str):
    wb = openpyxl.Workbook()

    # ========== 商品マスタシート ==========
    ws_prod = wb.active
    ws_prod.title = "商品マスタ"
    ws_prod.sheet_properties.tabColor = "70AD47"

    prod_headers = ["商品ID", "商品名", "カテゴリ", "単価"]
    for col, h in enumerate(prod_headers, 1):
        ws_prod.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_prod, 1, 1, len(prod_headers))

    for i, p in enumerate(data.PRODUCTS, 2):
        ws_prod.cell(row=i, column=1, value=p["id"]).font = styles.FONT_NORMAL
        ws_prod.cell(row=i, column=1).border = styles.THIN_BORDER
        ws_prod.cell(row=i, column=2, value=p["name"]).font = styles.FONT_NORMAL
        ws_prod.cell(row=i, column=2).border = styles.THIN_BORDER
        ws_prod.cell(row=i, column=3, value=p["category"]).font = styles.FONT_NORMAL
        ws_prod.cell(row=i, column=3).border = styles.THIN_BORDER
        c = ws_prod.cell(row=i, column=4, value=p["unit_price"])
        c.font = styles.FONT_NORMAL
        c.border = styles.THIN_BORDER
        c.number_format = styles.NUMBER_FORMAT_YEN

    for i, w in enumerate([12, 26, 16, 14], 1):
        ws_prod.column_dimensions[get_column_letter(i)].width = w

    # ========== 顧客マスタシート ==========
    ws_cust = wb.create_sheet("顧客マスタ")
    ws_cust.sheet_properties.tabColor = "70AD47"

    cust_headers = ["顧客ID", "顧客名", "地域", "業種"]
    for col, h in enumerate(cust_headers, 1):
        ws_cust.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_cust, 1, 1, len(cust_headers))

    for i, c in enumerate(data.CUSTOMERS, 2):
        ws_cust.cell(row=i, column=1, value=c["id"]).font = styles.FONT_NORMAL
        ws_cust.cell(row=i, column=1).border = styles.THIN_BORDER
        ws_cust.cell(row=i, column=2, value=c["name"]).font = styles.FONT_NORMAL
        ws_cust.cell(row=i, column=2).border = styles.THIN_BORDER
        ws_cust.cell(row=i, column=3, value=c["region"]).font = styles.FONT_NORMAL
        ws_cust.cell(row=i, column=3).border = styles.THIN_BORDER
        ws_cust.cell(row=i, column=4, value=c["industry"]).font = styles.FONT_NORMAL
        ws_cust.cell(row=i, column=4).border = styles.THIN_BORDER

    for i, w in enumerate([12, 30, 10, 10], 1):
        ws_cust.column_dimensions[get_column_letter(i)].width = w

    # ========== 売上データ (商品名なし — VLOOKUPで引く) ==========
    ws_sales = wb.create_sheet("売上データ")
    ws_sales.sheet_properties.tabColor = "70AD47"

    sale_headers = ["売上ID", "日付", "顧客ID", "商品ID", "数量", "金額", "担当者"]
    for col, h in enumerate(sale_headers, 1):
        ws_sales.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_sales, 1, 1, len(sale_headers))

    for i, s in enumerate(data.SALES[:50], 2):  # 50件に絞る
        row_data = [s["id"], s["date"], s["customer_id"], s["product_id"],
                    s["quantity"], s["amount"], s["rep"]]
        for col, val in enumerate(row_data, 1):
            cell = ws_sales.cell(row=i, column=col, value=val)
            cell.font = styles.FONT_NORMAL
            cell.border = styles.THIN_BORDER
            if col == 2:
                cell.number_format = "YYYY/MM/DD"
            elif col == 6:
                cell.number_format = styles.NUMBER_FORMAT_YEN

    for i, w in enumerate([10, 12, 10, 10, 8, 14, 14], 1):
        ws_sales.column_dimensions[get_column_letter(i)].width = w

    # ========== 問題シート ==========
    ws = wb.create_sheet("問題")
    styles.setup_problem_sheet(ws, "Lesson 2: VLOOKUP / INDEX-MATCH")
    ws.cell(row=3, column=1,
            value="各シートを参照して、VLOOKUPまたはINDEX-MATCH関数で回答してください。").font = styles.FONT_NORMAL

    # 期待値計算
    sale_0 = data.SALES[0]  # 最初の売上
    sale_4 = data.SALES[4]
    prod_p003 = data.get_product_by_id("P003")
    cust_c001 = data.get_customer_by_id("C001")
    cust_c010 = data.get_customer_by_id("C010")
    prod_p006 = data.get_product_by_id("P006")
    prod_p001 = data.get_product_by_id("P001")
    cust_c005 = data.get_customer_by_id("C005")

    questions = [
        ("Q1", "売上データの1行目（S0001）の商品IDから、商品名をVLOOKUPで求めてください",
         f"商品マスタの商品ID→商品名", sale_0["product_name"], False),
        ("Q2", "商品ID「P003」の単価をVLOOKUPで求めてください",
         "商品マスタから検索", prod_p003["unit_price"], True),
        ("Q3", "顧客ID「C001」の顧客名をVLOOKUPで求めてください",
         "顧客マスタから検索", cust_c001["name"], False),
        ("Q4", "顧客ID「C010」の地域をVLOOKUPで求めてください",
         "顧客マスタから検索", cust_c010["region"], False),
        ("Q5", "商品ID「P006」のカテゴリをINDEX-MATCHで求めてください",
         "=INDEX(範囲,MATCH(検索値,検索範囲,0))", prod_p006["category"], False),
        ("Q6", "商品ID「P001」の単価をINDEX-MATCHで求めてください",
         "=INDEX(範囲,MATCH(検索値,検索範囲,0))", prod_p001["unit_price"], True),
        ("Q7", "顧客ID「C005」の業種をINDEX-MATCHで求めてください",
         "=INDEX(範囲,MATCH(検索値,検索範囲,0))", cust_c005["industry"], False),
        ("Q8", "売上データ5行目（S0005）の顧客IDから、顧客名をVLOOKUPで求めてください",
         "売上→顧客マスタの2段階参照", data.get_customer_by_id(sale_4["customer_id"])["name"], False),
    ]

    row = 5
    for col, h in enumerate(["No.", "問題", "ヒント", "あなたの回答", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 5)

    answer_col = 4
    grading_col = 5
    start_row = row + 1

    for i, (qno, question, hint, expected, is_num) in enumerate(questions):
        r = start_row + i
        ws.cell(row=r, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=r, column=1).border = styles.THIN_BORDER
        ws.cell(row=r, column=2, value=question).font = styles.FONT_NORMAL
        ws.cell(row=r, column=2).border = styles.THIN_BORDER
        ws.cell(row=r, column=3, value=hint).font = styles.FONT_NORMAL
        ws.cell(row=r, column=3).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, r, answer_col)

        answer_ref = f"D{r}"
        if is_num:
            styles.add_grading_cell(ws, r, grading_col, answer_ref, expected)
        else:
            # 文字列比較: 解答シートのセル参照
            ws_ans_ref = f"'解答'!B{i + 2}"
            styles.add_grading_formula_cell(ws, r, grading_col, answer_ref, ws_ans_ref)
            # Override with text comparison for strings
            cell = ws.cell(row=r, column=grading_col)
            cell.value = f"=IF({answer_ref}=\"\",\"\",IF({answer_ref}='解答'!B{i+2},\"✓ 正解\",\"✗ 不正解\"))"

    end_row = start_row + len(questions) - 1
    styles.apply_grading_format(ws, grading_col, start_row, end_row)
    styles.add_score_summary(ws, end_row + 2, 1, grading_col, start_row, end_row, len(questions))

    ws.cell(row=end_row + 4, column=1,
            value="💡 全問回答後、E列のフォント色を黒に変更すると採点結果が見えます。").font = styles.FONT_NORMAL

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 14

    # ========== 解答シート (非表示) ==========
    ws_ans = styles.create_answer_sheet(wb, ws)
    ws_ans.cell(row=1, column=1, value="No.").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=2, value="期待値").font = styles.FONT_BOLD
    for i, (qno, _, _, expected, _) in enumerate(questions):
        ws_ans.cell(row=i + 2, column=1, value=qno)
        ws_ans.cell(row=i + 2, column=2, value=expected)

    wb.move_sheet("問題", offset=-3)

    path = f"{output_dir}/excel/Lesson2_VLOOKUP_INDEX-MATCH.xlsx"
    wb.save(path)
    print(f"  ✓ {path}")
    return path
