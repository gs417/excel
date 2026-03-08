"""Excel Lesson 4: ピボットテーブル + チャート — 5タスク"""

import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from collections import defaultdict
from . import data, styles


def generate(output_dir: str):
    wb = openpyxl.Workbook()

    # ========== 売上データシート ==========
    ws_data = wb.active
    ws_data.title = "売上データ"
    ws_data.sheet_properties.tabColor = "70AD47"

    headers = ["売上ID", "日付", "商品名", "カテゴリ", "単価", "数量", "金額", "担当者", "地域", "月"]
    for col, h in enumerate(headers, 1):
        ws_data.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_data, 1, 1, len(headers))

    for i, s in enumerate(data.SALES, 2):
        row_data = [
            s["id"], s["date"], s["product_name"], s["category"],
            s["unit_price"], s["quantity"], s["amount"],
            s["rep"], s["region"], s["month"],
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws_data.cell(row=i, column=col, value=val)
            cell.font = styles.FONT_NORMAL
            cell.border = styles.THIN_BORDER
            if col == 2:
                cell.number_format = "YYYY/MM/DD"
            elif col in (5, 7):
                cell.number_format = styles.NUMBER_FORMAT_YEN

    col_widths = [10, 12, 24, 14, 14, 8, 14, 14, 10, 10]
    for i, w in enumerate(col_widths, 1):
        ws_data.column_dimensions[get_column_letter(i)].width = w

    # ========== 集計値の事前計算 ==========
    # タスク1: 地域別売上合計
    region_totals = defaultdict(int)
    for s in data.SALES:
        region_totals[s["region"]] += s["amount"]

    # タスク2: カテゴリ別売上合計
    cat_totals = defaultdict(int)
    for s in data.SALES:
        cat_totals[s["category"]] += s["amount"]

    # タスク3: 担当者別売上件数
    rep_counts = defaultdict(int)
    for s in data.SALES:
        rep_counts[s["rep"]] += 1

    # タスク4: 月別売上合計 (上位3ヶ月)
    month_totals = defaultdict(int)
    for s in data.SALES:
        month_totals[s["month"]] += s["amount"]
    top3_months = sorted(month_totals.items(), key=lambda x: x[1], reverse=True)[:3]

    # タスク5: 地域×カテゴリ クロス集計 (東京×牛乳)
    tokyo_milk = sum(s["amount"] for s in data.SALES if s["region"] == "東京" and s["category"] == "牛乳")

    # ========== 問題シート ==========
    ws = wb.create_sheet("問題")
    styles.setup_problem_sheet(ws, "Lesson 4: ピボットテーブル + チャート")
    ws.cell(row=3, column=1,
            value="売上データからピボットテーブルを手動作成し、集計結果の値を入力してください。").font = styles.FONT_NORMAL
    ws.cell(row=4, column=1,
            value="※ 実務ではExcelのピボットテーブル機能を使いますが、ここでは理解のため手動集計します。").font = styles.FONT_NORMAL

    row = 6

    # --- タスク1: 地域別売上合計 ---
    ws.cell(row=row, column=1, value="タスク1: 地域別の売上合計を入力してください").font = styles.FONT_BOLD
    row += 1
    for col, h in enumerate(["地域", "売上合計（あなたの回答）", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 3)
    row += 1
    task1_start = row
    for region in data.REGIONS:
        ws.cell(row=row, column=1, value=region).font = styles.FONT_NORMAL
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, row, 2)
        styles.add_grading_cell(ws, row, 3, f"B{row}", region_totals[region])
        row += 1
    task1_end = row - 1

    row += 1

    # --- タスク2: カテゴリ別売上合計 ---
    ws.cell(row=row, column=1, value="タスク2: カテゴリ別の売上合計を入力してください").font = styles.FONT_BOLD
    row += 1
    for col, h in enumerate(["カテゴリ", "売上合計（あなたの回答）", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 3)
    row += 1
    task2_start = row
    categories = ["牛乳", "ヨーグルト", "チーズ", "バター・生クリーム"]
    for cat in categories:
        ws.cell(row=row, column=1, value=cat).font = styles.FONT_NORMAL
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, row, 2)
        styles.add_grading_cell(ws, row, 3, f"B{row}", cat_totals[cat])
        row += 1
    task2_end = row - 1

    row += 1

    # --- タスク3: 担当者別件数 ---
    ws.cell(row=row, column=1, value="タスク3: 担当者別の売上件数を入力してください").font = styles.FONT_BOLD
    row += 1
    for col, h in enumerate(["担当者", "件数（あなたの回答）", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 3)
    row += 1
    task3_start = row
    for rep in data.SALES_REPS:
        ws.cell(row=row, column=1, value=rep["name"]).font = styles.FONT_NORMAL
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, row, 2)
        styles.add_grading_cell(ws, row, 3, f"B{row}", rep_counts[rep["name"]])
        row += 1
    task3_end = row - 1

    row += 1

    # --- タスク4: 売上TOP3の月 ---
    ws.cell(row=row, column=1, value="タスク4: 月別売上合計のTOP3の月と金額を入力してください").font = styles.FONT_BOLD
    row += 1
    for col, h in enumerate(["順位", "月（YYYY-MM）", "売上合計", "採点"], 1):
        ws.cell(row=row, column=col, value=h)
    styles.style_header_row(ws, row, 1, 4)
    row += 1
    task4_start = row
    for rank, (month, amount) in enumerate(top3_months, 1):
        ws.cell(row=row, column=1, value=f"{rank}位").font = styles.FONT_NORMAL
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, row, 2)
        styles.mark_answer_cell(ws, row, 3)
        # 金額の採点
        styles.add_grading_cell(ws, row, 4, f"C{row}", amount)
        row += 1
    task4_end = row - 1

    row += 1

    # --- タスク5: クロス集計 ---
    ws.cell(row=row, column=1,
            value="タスク5: 東京地域 × 牛乳カテゴリ の売上合計を入力してください").font = styles.FONT_BOLD
    row += 1
    ws.cell(row=row, column=1, value="東京 × 牛乳").font = styles.FONT_NORMAL
    ws.cell(row=row, column=1).border = styles.THIN_BORDER
    styles.mark_answer_cell(ws, row, 2)
    styles.add_grading_cell(ws, row, 3, f"B{row}", tokyo_milk)
    task5_row = row

    # 条件付き書式を全採点列に適用
    styles.apply_grading_format(ws, 3, task1_start, task1_end)
    styles.apply_grading_format(ws, 3, task2_start, task2_end)
    styles.apply_grading_format(ws, 3, task3_start, task3_end)
    styles.apply_grading_format(ws, 4, task4_start, task4_end)
    styles.apply_grading_format(ws, 3, task5_row, task5_row)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14

    # ========== 模範チャートシート ==========
    ws_chart = wb.create_sheet("模範チャート")
    ws_chart.sheet_properties.tabColor = "FFC000"

    # 地域別データ書き込み
    ws_chart.cell(row=1, column=1, value="地域別売上合計").font = styles.FONT_TITLE
    ws_chart.cell(row=2, column=1, value="地域").font = styles.FONT_BOLD
    ws_chart.cell(row=2, column=2, value="売上合計").font = styles.FONT_BOLD
    for i, region in enumerate(data.REGIONS, 3):
        ws_chart.cell(row=i, column=1, value=region)
        ws_chart.cell(row=i, column=2, value=region_totals[region])
        ws_chart.cell(row=i, column=2).number_format = styles.NUMBER_FORMAT_YEN

    # 棒グラフ: 地域別
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.title = "地域別売上合計"
    bar_chart.y_axis.title = "金額"
    bar_chart.x_axis.title = "地域"
    bar_data = Reference(ws_chart, min_col=2, min_row=2, max_row=7)
    bar_cats = Reference(ws_chart, min_col=1, min_row=3, max_row=7)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    bar_chart.width = 18
    bar_chart.height = 12
    ws_chart.add_chart(bar_chart, "D2")

    # カテゴリ別データ
    ws_chart.cell(row=10, column=1, value="カテゴリ別売上構成").font = styles.FONT_TITLE
    ws_chart.cell(row=11, column=1, value="カテゴリ").font = styles.FONT_BOLD
    ws_chart.cell(row=11, column=2, value="売上合計").font = styles.FONT_BOLD
    for i, cat in enumerate(categories, 12):
        ws_chart.cell(row=i, column=1, value=cat)
        ws_chart.cell(row=i, column=2, value=cat_totals[cat])

    # 円グラフ: カテゴリ別
    pie_chart = PieChart()
    pie_chart.title = "カテゴリ別売上構成比"
    pie_data = Reference(ws_chart, min_col=2, min_row=11, max_row=15)
    pie_cats = Reference(ws_chart, min_col=1, min_row=12, max_row=15)
    pie_chart.add_data(pie_data, titles_from_data=True)
    pie_chart.set_categories(pie_cats)
    pie_chart.width = 16
    pie_chart.height = 12
    ws_chart.add_chart(pie_chart, "D18")

    # ========== 解答シート ==========
    ws_ans = styles.create_answer_sheet(wb, ws)
    ws_ans.cell(row=1, column=1, value="タスク").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=2, value="項目").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=3, value="期待値").font = styles.FONT_BOLD
    r = 2
    for region in data.REGIONS:
        ws_ans.cell(row=r, column=1, value="タスク1")
        ws_ans.cell(row=r, column=2, value=region)
        ws_ans.cell(row=r, column=3, value=region_totals[region])
        r += 1
    for cat in categories:
        ws_ans.cell(row=r, column=1, value="タスク2")
        ws_ans.cell(row=r, column=2, value=cat)
        ws_ans.cell(row=r, column=3, value=cat_totals[cat])
        r += 1
    for rep in data.SALES_REPS:
        ws_ans.cell(row=r, column=1, value="タスク3")
        ws_ans.cell(row=r, column=2, value=rep["name"])
        ws_ans.cell(row=r, column=3, value=rep_counts[rep["name"]])
        r += 1

    wb.move_sheet("問題", offset=-2)

    path = f"{output_dir}/excel/Lesson4_ピボットテーブル_チャート.xlsx"
    wb.save(path)
    print(f"  ✓ {path}")
    return path
