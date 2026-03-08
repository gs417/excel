"""Excel Lesson 5: KPIダッシュボード（総合） — 15問"""

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from collections import defaultdict
from . import data, styles


def generate(output_dir: str):
    wb = openpyxl.Workbook()

    # ========== 売上データシート ==========
    ws_data = wb.active
    ws_data.title = "売上データ"
    ws_data.sheet_properties.tabColor = "70AD47"

    headers = ["売上ID", "日付", "顧客ID", "商品名", "カテゴリ", "単価", "数量", "金額", "担当者", "地域", "月"]
    for col, h in enumerate(headers, 1):
        ws_data.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_data, 1, 1, len(headers))

    for i, s in enumerate(data.SALES, 2):
        row_data = [
            s["id"], s["date"], s["customer_id"], s["product_name"],
            s["category"], s["unit_price"], s["quantity"], s["amount"],
            s["rep"], s["region"], s["month"],
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws_data.cell(row=i, column=col, value=val)
            cell.font = styles.FONT_NORMAL
            cell.border = styles.THIN_BORDER
            if col == 2:
                cell.number_format = "YYYY/MM/DD"
            elif col in (6, 8):
                cell.number_format = styles.NUMBER_FORMAT_YEN

    # ========== 目標シート ==========
    ws_target = wb.create_sheet("目標")
    ws_target.sheet_properties.tabColor = "ED7D31"
    target_headers = ["担当者", "部署", "年間目標"]
    for col, h in enumerate(target_headers, 1):
        ws_target.cell(row=1, column=col, value=h)
    styles.style_header_row(ws_target, 1, 1, len(target_headers))

    for i, rep in enumerate(data.SALES_REPS, 2):
        ws_target.cell(row=i, column=1, value=rep["name"]).font = styles.FONT_NORMAL
        ws_target.cell(row=i, column=1).border = styles.THIN_BORDER
        ws_target.cell(row=i, column=2, value=rep["department"]).font = styles.FONT_NORMAL
        ws_target.cell(row=i, column=2).border = styles.THIN_BORDER
        c = ws_target.cell(row=i, column=3, value=rep["target"])
        c.font = styles.FONT_NORMAL
        c.border = styles.THIN_BORDER
        c.number_format = styles.NUMBER_FORMAT_YEN

    for i, w in enumerate([14, 14, 16], 1):
        ws_target.column_dimensions[get_column_letter(i)].width = w

    # ========== 事前計算 ==========
    total_amount = sum(s["amount"] for s in data.SALES)
    total_deals = len(data.SALES)
    avg_deal = total_amount / total_deals
    total_target = sum(r["target"] for r in data.SALES_REPS)
    achievement_rate = total_amount / total_target

    unique_customers = len(set(s["customer_id"] for s in data.SALES))

    rep_amounts = defaultdict(int)
    for s in data.SALES:
        rep_amounts[s["rep"]] += s["amount"]

    top_rep = max(rep_amounts, key=rep_amounts.get)
    top_rep_amount = rep_amounts[top_rep]

    month_amounts = defaultdict(int)
    month_counts = defaultdict(int)
    for s in data.SALES:
        month_amounts[s["month"]] += s["amount"]
        month_counts[s["month"]] += 1

    best_month = max(month_amounts, key=month_amounts.get)
    best_month_amount = month_amounts[best_month]

    region_amounts = defaultdict(int)
    for s in data.SALES:
        region_amounts[s["region"]] += s["amount"]
    top_region = max(region_amounts, key=region_amounts.get)
    top_region_amount = region_amounts[top_region]

    cat_amounts = defaultdict(int)
    for s in data.SALES:
        cat_amounts[s["category"]] += s["amount"]
    top_cat = max(cat_amounts, key=cat_amounts.get)

    # 達成者数
    achievers = sum(1 for r in data.SALES_REPS if rep_amounts[r["name"]] >= r["target"])

    # Q2H (上半期: 2025-04 ~ 2025-09)
    h1_amount = sum(s["amount"] for s in data.SALES if s["month"] <= "2025-09")
    h2_amount = total_amount - h1_amount

    # 平均月間売上
    months_count = len(month_amounts)
    avg_monthly = total_amount / months_count if months_count else 0

    # ========== ダッシュボードシート ==========
    ws = wb.create_sheet("ダッシュボード")
    styles.setup_problem_sheet(ws, "Lesson 5: KPIダッシュボード（総合演習）")
    ws.cell(row=3, column=1,
            value="半完成ダッシュボードの黄色セルに数式や値を入力して完成させてください。").font = styles.FONT_NORMAL

    # --- セクション1: 全体KPI ---
    row = 5
    ws.cell(row=row, column=1, value="■ 全体KPI").font = styles.FONT_TITLE
    ws.cell(row=row, column=1).fill = styles.FILL_LIGHT_BLUE
    for c in range(1, 6):
        ws.cell(row=row, column=c).fill = styles.FILL_LIGHT_BLUE

    row += 1
    kpi_items = [
        ("Q1", "売上合計", total_amount, styles.NUMBER_FORMAT_YEN),
        ("Q2", "取引件数", total_deals, styles.NUMBER_FORMAT_INT),
        ("Q3", "平均取引額", avg_deal, styles.NUMBER_FORMAT_YEN),
        ("Q4", "目標達成率", achievement_rate, styles.NUMBER_FORMAT_PCT),
        ("Q5", "取引顧客数（ユニーク）", unique_customers, styles.NUMBER_FORMAT_INT),
    ]

    q_start_row = row
    grading_col = 4  # D列
    answer_col = 3   # C列

    for qno, label, expected, fmt in kpi_items:
        ws.cell(row=row, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        ws.cell(row=row, column=2, value=label).font = styles.FONT_NORMAL
        ws.cell(row=row, column=2).border = styles.THIN_BORDER
        ac = styles.mark_answer_cell(ws, row, answer_col)
        ac.number_format = fmt
        styles.add_grading_cell(ws, row, grading_col, f"C{row}", expected)
        row += 1

    row += 1

    # --- セクション2: 担当者分析 ---
    ws.cell(row=row, column=1, value="■ 担当者分析").font = styles.FONT_TITLE
    ws.cell(row=row, column=1).fill = styles.FILL_LIGHT_BLUE
    for c in range(1, 6):
        ws.cell(row=row, column=c).fill = styles.FILL_LIGHT_BLUE
    row += 1

    rep_items = [
        ("Q6", "売上No.1の担当者名", top_rep, None, False),
        ("Q7", "その担当者の売上合計", top_rep_amount, styles.NUMBER_FORMAT_YEN, True),
        ("Q8", "目標達成者数（5名中）", achievers, styles.NUMBER_FORMAT_INT, True),
    ]

    for qno, label, expected, fmt, is_num in rep_items:
        ws.cell(row=row, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        ws.cell(row=row, column=2, value=label).font = styles.FONT_NORMAL
        ws.cell(row=row, column=2).border = styles.THIN_BORDER
        ac = styles.mark_answer_cell(ws, row, answer_col)
        if fmt:
            ac.number_format = fmt
        if is_num:
            styles.add_grading_cell(ws, row, grading_col, f"C{row}", expected)
        else:
            cell = ws.cell(row=row, column=grading_col)
            cell.value = f'=IF(C{row}="","",IF(C{row}=解答!B{qno[1:]},"✓ 正解","✗ 不正解"))'
            cell.font = styles.FONT_HIDDEN
            cell.alignment = styles.ALIGN_CENTER
            cell.border = styles.THIN_BORDER
        row += 1

    row += 1

    # --- セクション3: 時系列分析 ---
    ws.cell(row=row, column=1, value="■ 時系列分析").font = styles.FONT_TITLE
    ws.cell(row=row, column=1).fill = styles.FILL_LIGHT_BLUE
    for c in range(1, 6):
        ws.cell(row=row, column=c).fill = styles.FILL_LIGHT_BLUE
    row += 1

    time_items = [
        ("Q9", "売上最大月（YYYY-MM形式）", best_month, None, False),
        ("Q10", "その月の売上合計", best_month_amount, styles.NUMBER_FORMAT_YEN, True),
        ("Q11", "上半期(4月~9月)売上合計", h1_amount, styles.NUMBER_FORMAT_YEN, True),
        ("Q12", "下半期(10月~3月)売上合計", h2_amount, styles.NUMBER_FORMAT_YEN, True),
        ("Q13", "平均月間売上", avg_monthly, styles.NUMBER_FORMAT_YEN, True),
    ]

    for qno, label, expected, fmt, is_num in time_items:
        ws.cell(row=row, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        ws.cell(row=row, column=2, value=label).font = styles.FONT_NORMAL
        ws.cell(row=row, column=2).border = styles.THIN_BORDER
        ac = styles.mark_answer_cell(ws, row, answer_col)
        if fmt:
            ac.number_format = fmt
        if is_num:
            styles.add_grading_cell(ws, row, grading_col, f"C{row}", expected)
        else:
            cell = ws.cell(row=row, column=grading_col)
            cell.value = f'=IF(C{row}="","",IF(C{row}=解答!B{qno[1:]},"✓ 正解","✗ 不正解"))'
            cell.font = styles.FONT_HIDDEN
            cell.alignment = styles.ALIGN_CENTER
            cell.border = styles.THIN_BORDER
        row += 1

    row += 1

    # --- セクション4: エリア・カテゴリ ---
    ws.cell(row=row, column=1, value="■ エリア・カテゴリ分析").font = styles.FONT_TITLE
    ws.cell(row=row, column=1).fill = styles.FILL_LIGHT_BLUE
    for c in range(1, 6):
        ws.cell(row=row, column=c).fill = styles.FILL_LIGHT_BLUE
    row += 1

    area_items = [
        ("Q14", "売上No.1の地域", top_region, None, False),
        ("Q15", "売上No.1のカテゴリ", top_cat, None, False),
    ]

    for qno, label, expected, fmt, is_num in area_items:
        ws.cell(row=row, column=1, value=qno).font = styles.FONT_BOLD
        ws.cell(row=row, column=1).border = styles.THIN_BORDER
        ws.cell(row=row, column=2, value=label).font = styles.FONT_NORMAL
        ws.cell(row=row, column=2).border = styles.THIN_BORDER
        styles.mark_answer_cell(ws, row, answer_col)
        cell = ws.cell(row=row, column=grading_col)
        cell.value = f'=IF(C{row}="","",IF(C{row}=解答!B{qno[1:]},"✓ 正解","✗ 不正解"))'
        cell.font = styles.FONT_HIDDEN
        cell.alignment = styles.ALIGN_CENTER
        cell.border = styles.THIN_BORDER
        row += 1

    q_end_row = row - 1

    # 条件付き書式
    styles.apply_grading_format(ws, grading_col, q_start_row, q_end_row)

    # スコアサマリー
    row += 1
    styles.add_score_summary(ws, row, 1, grading_col, q_start_row, q_end_row, 15)

    ws.cell(row=row + 2, column=1,
            value="💡 全問回答後、D列のフォント色を黒に変更すると採点結果が見えます。").font = styles.FONT_NORMAL

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14

    # ========== 模範チャートシート ==========
    ws_chart = wb.create_sheet("模範チャート")

    # 月別売上推移データ
    ws_chart.cell(row=1, column=1, value="月別売上推移").font = styles.FONT_TITLE
    sorted_months = sorted(month_amounts.keys())
    ws_chart.cell(row=2, column=1, value="月").font = styles.FONT_BOLD
    ws_chart.cell(row=2, column=2, value="売上合計").font = styles.FONT_BOLD
    for i, m in enumerate(sorted_months, 3):
        ws_chart.cell(row=i, column=1, value=m)
        ws_chart.cell(row=i, column=2, value=month_amounts[m])

    line_chart = LineChart()
    line_chart.title = "月別売上推移"
    line_chart.y_axis.title = "金額"
    line_data = Reference(ws_chart, min_col=2, min_row=2, max_row=2 + len(sorted_months))
    line_cats = Reference(ws_chart, min_col=1, min_row=3, max_row=2 + len(sorted_months))
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)
    line_chart.width = 24
    line_chart.height = 14
    ws_chart.add_chart(line_chart, "D2")

    # 担当者別棒グラフ
    chart_row = 4 + len(sorted_months)
    ws_chart.cell(row=chart_row, column=1, value="担当者別売上").font = styles.FONT_TITLE
    ws_chart.cell(row=chart_row + 1, column=1, value="担当者").font = styles.FONT_BOLD
    ws_chart.cell(row=chart_row + 1, column=2, value="売上合計").font = styles.FONT_BOLD
    ws_chart.cell(row=chart_row + 1, column=3, value="目標").font = styles.FONT_BOLD
    for i, rep in enumerate(data.SALES_REPS):
        r = chart_row + 2 + i
        ws_chart.cell(row=r, column=1, value=rep["name"])
        ws_chart.cell(row=r, column=2, value=rep_amounts[rep["name"]])
        ws_chart.cell(row=r, column=3, value=rep["target"])

    bar = BarChart()
    bar.type = "col"
    bar.title = "担当者別 実績 vs 目標"
    bar_data = Reference(ws_chart, min_col=2, max_col=3, min_row=chart_row + 1,
                         max_row=chart_row + 1 + len(data.SALES_REPS))
    bar_cats = Reference(ws_chart, min_col=1, min_row=chart_row + 2,
                         max_row=chart_row + 1 + len(data.SALES_REPS))
    bar.add_data(bar_data, titles_from_data=True)
    bar.set_categories(bar_cats)
    bar.width = 24
    bar.height = 14
    ws_chart.add_chart(bar, "D20")

    # ========== 解答シート ==========
    ws_ans = styles.create_answer_sheet(wb, ws)
    ws_ans.cell(row=1, column=1, value="No.").font = styles.FONT_BOLD
    ws_ans.cell(row=1, column=2, value="期待値").font = styles.FONT_BOLD

    all_expected = [
        total_amount, total_deals, avg_deal, achievement_rate, unique_customers,
        top_rep, top_rep_amount, achievers,
        best_month, best_month_amount, h1_amount, h2_amount, avg_monthly,
        top_region, top_cat,
    ]
    for i, val in enumerate(all_expected, 1):
        ws_ans.cell(row=i + 1, column=1, value=i)
        ws_ans.cell(row=i + 1, column=2, value=val)

    wb.move_sheet("ダッシュボード", offset=-3)

    path = f"{output_dir}/excel/Lesson5_KPIダッシュボード.xlsx"
    wb.save(path)
    print(f"  ✓ {path}")
    return path
